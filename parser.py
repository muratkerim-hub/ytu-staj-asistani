import re
try:
    import pymupdf as fitz  # PyMuPDF kütüphanesi (Modern import)
except ImportError:
    import fitz
from docx import Document  # python-docx kütüphanesi (Word okumak için)

# 1. Türkçe ve İngilizce CV bölüm başlıkları sözlüğümüz
SECTION_KEYWORDS = {
    "iletisim": [
        "iletişim", "iletisim", "contact", "kişisel bilgiler", "personal info",
        "iletişim bilgileri", "iletisim bilgileri", "contact info", "contact details"
    ],
    "ozet": [
        "özet", "ozet", "hakkımda", "hakkimda", "summary", "profile", "about me",
        "kariyer hedefi", "profesyonel özet", "özgeçmiş özeti", "professional summary",
        "özet / profil", "özet ve profil", "profil ve özet"
    ],
    "egitim": [
        "eğitim", "egitim", "education", "akademik", "academic", "eğitim bilgileri",
        "egitim bilgileri", "eğitim durumu", "educational background"
    ],
    "deneyim": [
        "deneyim", "deneyimler", "deneyimlerim", "iş deneyimi", "is deneyimi",
        "iş tecrübesi", "is tecrubesi", "staj", "stajlar", "staj deneyimi",
        "deneyim & stajlar", "deneyim ve stajlar", "iş deneyimi ve stajlar",
        "experience", "work experience", "employment", "internship", "internships", "professional experience"
    ],
    "projeler": [
        "projeler", "projelerim", "proje", "projects", "personal projects",
        "akademik projeler", "seçilmiş projeler", "projeler ve çalışmalar", "key projects"
    ],
    "yetenekler": [
        "yetenekler", "yeteneklerim", "beceriler", "yetkinlikler", "teknik beceriler",
        "teknik yetenekler", "teknik yetenekler & araçlar", "teknik yetenekler ve araçlar",
        "skills", "technical skills", "technologies", "araçlar ve yetenekler"
    ],
    "sertifikalar": [
        "sertifikalar", "sertifikalarım", "kurslar", "başarılar", "certificates",
        "certifications", "courses", "awards", "sertifikalar ve kurslar", "ödüller",
        "sertifikalar & diller", "sertifikalar ve diller"
    ],
    "diller": [
        "yabancı dil", "yabancı diller", "diller", "dil", "languages", "language", "dil yetkinlikleri"
    ]
}


def extract_text_from_pdf(file_path: str) -> dict:
    """PDF dosyasından metin ve sayfa bilgisi çıkarır."""
    try:
        doc = fitz.open(file_path)
        full_text = ""
        page_count = len(doc)
        has_images = False

        for page in doc:
            full_text += page.get_text("text") + "\n"
            if len(page.get_images()) > 0:
                has_images = True

        doc.close()
        clean_text = full_text.strip()

        # Taranmış görsel PDF kontrolü: Sayfa var ama içinden 30 karakter bile çıkmadıysa
        if len(clean_text) < 30:
            return {
                "success": False,
                "error": "⚠️ Bu PDF'ten metin okunamadı! Dosya taranmış bir resim (görsel PDF) olabilir. Lütfen metin seçilebilir bir PDF veya DOCX yükleyin.",
                "text": "",
                "page_count": page_count,
                "has_images": has_images
            }

        return {
            "success": True,
            "error": None,
            "text": clean_text,
            "page_count": page_count,
            "has_images": has_images
        }
    except Exception as e:
        return {"success": False, "error": f"PDF okuma hatası: {e}", "text": "", "page_count": 0, "has_images": False}


def extract_text_from_docx(file_path: str) -> dict:
    """DOCX (Word) dosyasından metin, liste maddeleri ve tablo kontrolü çıkarır."""
    try:
        doc = Document(file_path)
        paragraphs = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue

            # Word liste stilleri (List Bullet, List Number vb.) veya XML madde imi varsa
            # metnin başında açıkça madde imi yoksa '• ' ekleyerek ATS ayrıştırıcısının
            # maddeleri doğru ve tek tek tanımasını sağla.
            is_bullet = False
            if p.style and any(k in p.style.name.lower() for k in ["bullet", "list"]):
                is_bullet = True
            elif hasattr(p, "_element") and p._element.xpath("./w:pPr/w:numPr"):
                is_bullet = True

            if is_bullet and not re.match(r"^[\•\-\*\–\—\▪\▫\◦\►\>\d\.\)]+\s*", txt):
                paragraphs.append(f"• {txt}")
            else:
                paragraphs.append(txt)

        full_text = "\n".join(paragraphs)

        # ATS kontrolü için dosyada tablo var mı bilgisini de tutuyoruz
        has_tables = len(doc.tables) > 0

        if len(full_text.strip()) < 30:
            return {
                "success": False,
                "error": "⚠️ DOCX dosyası boş görünüyor veya metinler sadece tablo/görsel içine gömülmüş.",
                "text": "",
                "page_count": 1,
                "has_tables": has_tables
            }

        return {
            "success": True,
            "error": None,
            "text": full_text,
            "page_count": 1,
            "has_tables": has_tables
        }
    except Exception as e:
        return {"success": False, "error": f"DOCX okuma hatası: {e}", "text": "", "page_count": 0, "has_tables": False}


def turkish_lower(text: str) -> str:
    """Türkçe İ ve I harflerini doğru dönüştürerek küçük harfe çevirir."""
    return text.replace("İ", "i").replace("I", "ı").lower()


def normalize_header_text(text: str) -> str:
    """
    Bölüm başlıklarını karşılaştırmak için Türkçe harfleri, noktalama ve sembolleri normalize eder.
    Böylece 'EĞITIM', 'EĞİTİM', 'EGITIM', 'Özet / Profil', 'Deneyim & Stajlar' gibi
    tüm varyasyonlar firesiz ve kusursuz eşleşir.
    """
    if not text:
        return ""
    # Küçük harfe çevirirken hem Türkçe İ/I hem İngilizce I durumlarını eşitle
    t = text.replace("İ", "i").replace("I", "i").replace("ı", "i").lower()
    t = t.replace("ö", "o").replace("ü", "u").replace("ç", "c").replace("ş", "s").replace("ğ", "g")
    # & ve / sembollerini boşluk/ve ile normalize et
    t = t.replace("&", " ve ").replace("/", " ve ")
    # Fazla boşlukları ve baştaki/sondaki noktalama işaretlerini temizle
    t = re.sub(r"[\:\-\–\—\•\*\.\,\;]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


# Hızlı ve firesiz eşleşme için önceden normalize edilmiş anahtar kelimeler
NORMALIZED_SECTION_KEYWORDS = {
    section_key: [normalize_header_text(kw) for kw in keywords]
    for section_key, keywords in SECTION_KEYWORDS.items()
}


def detect_section_header(line: str) -> str | None:
    """
    Bir satırın bölüm başlığı (örn. 'EĞİTİM', 'Özet / Profil', 'Deneyim & Stajlar') olup olmadığını kontrol eder.
    Başlıksa bölüm adını döndürür, değilse None döndürür.
    """
    raw = line.strip()
    if not raw:
        return None

    # Bir başlık satırı genelde kısadır (en fazla 6 kelime ve 55 karakter)
    if len(raw) > 55 or len(raw.split()) > 6:
        return None

    norm = normalize_header_text(raw)
    if not norm:
        return None

    # 1. Tam eşleşme (normalize edilmiş anahtar kelimeler ile)
    for section_key, norm_keywords in NORMALIZED_SECTION_KEYWORDS.items():
        for kw in norm_keywords:
            if norm == kw:
                return section_key

    # 2. Bileşik başlık ve kısmi eşleşme desteği (tam kelime sınırları ile)
    for section_key, norm_keywords in NORMALIZED_SECTION_KEYWORDS.items():
        for kw in norm_keywords:
            if len(kw) >= 4:
                pattern = r"(?:^|\b)" + re.escape(kw) + r"(?:\b|$)"
                if re.search(pattern, norm):
                    return section_key

    return None


def parse_cv_sections(raw_text: str) -> dict:
    """Ham CV metnini satır satır okuyup bölümlere (sözlük yapısına) ayırır."""
    # Başlangıçta tüm bölümleri boş liste olarak hazırla
    sections = {key: [] for key in SECTION_KEYWORDS.keys()}
    sections["baslik_ve_kisisel"] = []  # En üstteki isim ve iletişim satırları için

    current_section = "baslik_ve_kisisel"

    for line in raw_text.splitlines():
        stripped_line = line.strip()
        if not stripped_line:
            continue  # Boş satırları atla

        # Bu satır yeni bir bölüm başlığı mı?
        detected = detect_section_header(stripped_line)
        if detected:
            current_section = detected
        else:
            # Başlık değilse, o an hangi bölümdeysek o bölümün listesine ekle
            sections[current_section].append(stripped_line)

    # Listeleri tek metin haline getir, boş kalan bölümleri temiz tut
    parsed_result = {}
    for sec, lines in sections.items():
        parsed_result[sec] = "\n".join(lines)

    return parsed_result


# --- TEST BLOĞU ---
# Bu dosyayı doğrudan çalıştırdığımızda örnek bir CV metni ile test eder
if __name__ == "__main__":
    ornek_cv = """
    Murat Yılmaz
    murat@ytu.edu.tr | github.com/murat | İstanbul

    ÖZET
    Yıldız Teknik Üniversitesi öğrencisiyim. Yapay zeka ve veri mühendisliği ile ilgileniyorum.

    EĞİTİM
    Yıldız Teknik Üniversitesi - Matematik Mühendisliği (2025 - Devam)

    PROJECTS
    Staj Asistanı: Python ve Streamlit kullanarak CV analiz aracı geliştirdim.

    SKILLS
    Python, C, SQL, Git, OOP
    """

    print("🔍 Örnek CV Ayrıştırılıyor...\n")
    bolumler = parse_cv_sections(ornek_cv)

    for bolum_adi, icerik in bolumler.items():
        if icerik:  # Sadece dolu olan bölümleri yazdır
            print(f"📌 [{bolum_adi.upper()}]")
            print(icerik)
            print("-" * 40)